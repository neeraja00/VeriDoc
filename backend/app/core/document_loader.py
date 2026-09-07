"""Document extraction and parsing module for PDF, DOCX, TXT, and Markdown."""
import io
import os
from pathlib import Path
from typing import List, Dict, Any, Tuple
from dataclasses import dataclass, field
import pypdf
from docx import Document as DocxDocument


class DocumentLoaderError(Exception):
    """Base exception for document loading errors."""
    pass


class UnsupportedFileTypeError(DocumentLoaderError):
    """Raised when an unsupported file extension or MIME type is encountered."""
    pass


class EmptyDocumentError(DocumentLoaderError):
    """Raised when an uploaded document has no extractable text."""
    pass


class CorruptedFileError(DocumentLoaderError):
    """Raised when a document cannot be parsed due to file corruption."""
    pass


@dataclass
class DocumentPage:
    """Represents an extracted page or structural unit from a document."""
    page_number: int
    text: str
    metadata: Dict[str, Any] = field(default_factory=dict)


class DocumentLoader:
    """Parses and extracts text and metadata from supported document formats."""

    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}

    @classmethod
    def load(cls, file_bytes: bytes, filename: str) -> Tuple[List[DocumentPage], Dict[str, Any]]:
        """
        Loads and parses a document from raw bytes.

        Returns:
            Tuple of (list of DocumentPage, document_metadata_dict)
        """
        ext = Path(filename).suffix.lower()
        if ext not in cls.SUPPORTED_EXTENSIONS:
            raise UnsupportedFileTypeError(
                f"Unsupported file format '{ext}'. Supported formats: {', '.join(sorted(cls.SUPPORTED_EXTENSIONS))}"
            )

        if not file_bytes or len(file_bytes.strip()) == 0:
            raise EmptyDocumentError(f"The file '{filename}' is empty.")

        try:
            if ext == ".pdf":
                pages = cls._load_pdf(file_bytes, filename)
            elif ext == ".docx":
                pages = cls._load_docx(file_bytes, filename)
            elif ext in {".txt", ".md"}:
                pages = cls._load_text(file_bytes, filename)
            else:
                raise UnsupportedFileTypeError(f"Unhandled file extension: {ext}")
        except DocumentLoaderError:
            raise
        except Exception as e:
            raise CorruptedFileError(f"Failed to parse document '{filename}': {str(e)}") from e

        # Filter out pages that are completely empty
        non_empty_pages = [p for p in pages if p.text.strip()]
        if not non_empty_pages:
            raise EmptyDocumentError(
                f"No extractable text found in '{filename}'. The file may contain only scanned images or blank pages."
            )

        total_chars = sum(len(p.text) for p in non_empty_pages)
        doc_meta = {
            "filename": filename,
            "file_type": ext.lstrip("."),
            "file_size_bytes": len(file_bytes),
            "total_pages": len(non_empty_pages),
            "total_characters": total_chars,
        }

        return non_empty_pages, doc_meta

    @staticmethod
    def _load_pdf(file_bytes: bytes, filename: str) -> List[DocumentPage]:
        """Extracts text from PDF page by page."""
        stream = io.BytesIO(file_bytes)
        try:
            reader = pypdf.PdfReader(stream)
        except Exception as e:
            raise CorruptedFileError(f"Invalid or corrupted PDF file '{filename}': {e}") from e

        if reader.is_encrypted:
            try:
                # Attempt empty password decrypt
                reader.decrypt("")
            except Exception:
                raise CorruptedFileError(f"PDF '{filename}' is encrypted/password-protected.")

        pages: List[DocumentPage] = []
        for idx, page in enumerate(reader.pages):
            try:
                extracted = page.extract_text() or ""
            except Exception:
                extracted = ""
            pages.append(
                DocumentPage(
                    page_number=idx + 1,
                    text=extracted.strip(),
                    metadata={"source_filename": filename, "page_number": idx + 1}
                )
            )
        return pages

    @staticmethod
    def _load_docx(file_bytes: bytes, filename: str) -> List[DocumentPage]:
        """Extracts text from DOCX document."""
        stream = io.BytesIO(file_bytes)
        try:
            doc = DocxDocument(stream)
        except Exception as e:
            raise CorruptedFileError(f"Invalid or corrupted Word document '{filename}': {e}") from e

        full_text_paragraphs = []
        for p in doc.paragraphs:
            if p.text and p.text.strip():
                full_text_paragraphs.append(p.text.strip())

        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    full_text_paragraphs.append(row_text)

        joined_text = "\n\n".join(full_text_paragraphs)
        return [
            DocumentPage(
                page_number=1,
                text=joined_text,
                metadata={"source_filename": filename, "page_number": 1}
            )
        ]

    @staticmethod
    def _load_text(file_bytes: bytes, filename: str) -> List[DocumentPage]:
        """Extracts text from plain text or markdown with encoding fallbacks."""
        text = ""
        for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
            try:
                text = file_bytes.decode(encoding)
                break
            except UnicodeDecodeError:
                continue

        if not text:
            text = file_bytes.decode("utf-8", errors="replace")

        return [
            DocumentPage(
                page_number=1,
                text=text.strip(),
                metadata={"source_filename": filename, "page_number": 1}
            )
        ]
